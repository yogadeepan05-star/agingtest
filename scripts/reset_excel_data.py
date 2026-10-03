import os
import shutil
from pathlib import Path
from openpyxl import Workbook

HEADERS = [
    'Serial Number', 'Device Registration Time', 'Registration Battery %',
    'H1 Battery %', 'H1 Timestamp', 'H2 Battery %', 'H2 Timestamp',
    'H3 Battery %', 'H3 Timestamp', 'H4 Battery %', 'H4 Timestamp',
    'Post-Aging Battery %', 'Post-Aging Timestamp', 'Final Status'
]

def create_fresh_workbook():
    book = Workbook()
    ws_devices = book.active
    ws_devices.title = 'Devices'
    ws_devices.append(HEADERS)
    ws_devices.freeze_panes = 'D2'
    ws_devices.auto_filter.ref = 'A1:N1'
    for col in ws_devices.columns:
        ws_devices.column_dimensions[col[0].column_letter].width = 28

    ws_workflow = book.create_sheet('Workflow')
    ws_workflow.append(['Serial Number', 'State JSON'])
    ws_workflow.sheet_state = 'hidden'

    ws_captures = book.create_sheet('Captures')
    ws_captures.append(['Token', 'Action', 'Serial', 'Issued At', 'Used', 'Revision'])
    ws_captures.sheet_state = 'hidden'

    return book

targets = [
    Path(r"C:\Users\DEEPAN V\Downloads\production-aging-test-photo-validation\production-aging-test\data\aging_test.xlsx"),
    Path(r"C:\Users\DEEPAN V\Downloads\production-aging-test\production-aging-test\data\aging_test.xlsx")
]

for target in targets:
    if target.parent.exists():
        if target.exists():
            backup = target.with_name("aging_test_backup.xlsx")
            shutil.copy2(target, backup)
            print(f"Backed up {target} to {backup}")
        wb = create_fresh_workbook()
        wb.save(target)
        wb.close()
        print(f"Successfully reset: {target}")
