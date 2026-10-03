import json
import os
import tempfile
import time
from pathlib import Path
from filelock import FileLock
from openpyxl import Workbook, load_workbook

HEADERS = ['Serial Number', 'Device Registration Time', 'Registration Battery %',
           'H1 Battery %', 'H1 Timestamp', 'H2 Battery %', 'H2 Timestamp',
           'H3 Battery %', 'H3 Timestamp', 'H4 Battery %', 'H4 Timestamp',
           'Post-Aging Battery %', 'Post-Aging Timestamp', 'Final Status']

NEW_HEADERS = [
    'H1 Issue', 'H1 Issue Categories', 'H1 Remarks',
    'H2 Issue', 'H2 Issue Categories', 'H2 Remarks',
    'H3 Issue', 'H3 Issue Categories', 'H3 Remarks',
    'H4 Issue', 'H4 Issue Categories', 'H4 Remarks',
    'Post-Aging Issue', 'Post-Aging Issue Categories', 'Post-Aging Remarks',
    'Long Press Power Test'
]

ALL_HEADERS = HEADERS + NEW_HEADERS


def sanitize_formula_injection(text: str | None) -> str | None:
    """Neutralize formula injection by prepending a single quote if string starts with formula characters."""
    if text is None:
        return None
    s = str(text).strip()
    if not s:
        return ''
    if s[0] in ('=', '+', '-', '@', '\t', '\r'):
        return "'" + s
    return s


def extract_observations_from_row(sheet, row: int):
    """Extract observations and power test result from row columns 15–30."""
    obs = {}
    for n in range(1, 5):
        issue_raw = sheet.cell(row, 15 + 3 * (n - 1)).value
        cat_raw = sheet.cell(row, 16 + 3 * (n - 1)).value
        rem_raw = sheet.cell(row, 17 + 3 * (n - 1)).value
        if issue_raw is not None and str(issue_raw).strip():
            issue_val = 'yes' if str(issue_raw).strip().lower() == 'yes' else 'no'
            cats = [c.strip() for c in str(cat_raw).split(',') if c.strip()] if cat_raw else []
            obs[f'h{n}'] = {
                'has_issue': issue_val,
                'categories': cats,
                'remarks': str(rem_raw) if rem_raw is not None else ''
            }
        else:
            obs[f'h{n}'] = None

    post_issue = sheet.cell(row, 27).value
    post_cat = sheet.cell(row, 28).value
    post_rem = sheet.cell(row, 29).value
    if post_issue is not None and str(post_issue).strip():
        issue_val = 'yes' if str(post_issue).strip().lower() == 'yes' else 'no'
        cats = [c.strip() for c in str(post_cat).split(',') if c.strip()] if post_cat else []
        obs['post'] = {
            'has_issue': issue_val,
            'categories': cats,
            'remarks': str(post_rem) if post_rem is not None else ''
        }
    else:
        obs['post'] = None

    power_test = sheet.cell(row, 30).value
    power_test_val = str(power_test).strip() if power_test is not None and str(power_test).strip() else None
    return obs, power_test_val

class Store:
    def __init__(self, path):
        self.path = Path(path)
        self.path.parent.mkdir(parents=True, exist_ok=True)
        self.lock = FileLock(str(path) + '.lock', timeout=10)

    @staticmethod
    def reconcile(book):
        active_serials = set()
        devices = book['Devices']
        for row in devices.iter_rows(min_row=2):
            val = row[0].value
            if val is not None and str(val).strip():
                active_serials.add(str(val).strip())

        if 'Workflow' in book.sheetnames:
            wf = book['Workflow']
            seen = set()
            for r in range(wf.max_row, 1, -1):
                val = wf.cell(r, 1).value
                s = str(val).strip() if val is not None else ''
                if not s or s not in active_serials or s in seen:
                    wf.delete_rows(r)
                else:
                    seen.add(s)

        if 'Captures' in book.sheetnames:
            cap = book['Captures']
            for r in range(cap.max_row, 1, -1):
                val = cap.cell(r, 3).value
                s = str(val).strip() if val is not None else ''
                if s and s not in active_serials:
                    cap.delete_rows(r)

    def transaction(self, operation, write=True):
        if self.path.is_symlink() or Path(str(self.path) + '.lock').is_symlink():
            raise ValueError('Symlinked workbook or lock is not supported')
        with self.lock:
            created_new = not self.path.exists()
            if not created_new:
                book = load_workbook(self.path)
                if 'Devices' not in book.sheetnames:
                    book.close()
                    raise ValueError('Workbook needs a Devices sheet with the documented headers; existing file preserved.')
                if [c.value for c in book['Devices'][1]][:14] != HEADERS:
                    book.close()
                    raise ValueError('Workbook headers do not match. Existing file preserved.')
                # Ensure new column headers exist without modifying existing data
                for idx, h_title in enumerate(NEW_HEADERS, start=15):
                    if book['Devices'].cell(1, idx).value is None:
                        book['Devices'].cell(1, idx, h_title)
            else:
                book = Workbook()
                book.active.title = 'Devices'
                book.active.append(ALL_HEADERS)
                book.active.freeze_panes = 'D2'
                book.active.auto_filter.ref = 'A1:AD1'
                for col in book.active.columns:
                    book.active.column_dimensions[col[0].column_letter].width = 28
            try:
                for name, headers in [('Workflow', ['Serial Number', 'State JSON']),
                                      ('Captures', ['Token', 'Action', 'Serial', 'Issued At', 'Used', 'Revision'])]:
                    if name not in book.sheetnames:
                        book.create_sheet(name).append(headers)
                        book[name].sheet_state = 'hidden'
                if book['Captures'].cell(1, 6).value != 'Revision':
                    book['Captures'].cell(1, 6, 'Revision')
                self.reconcile(book)
                result = operation(book)
                if write or created_new:
                    fd, temp = tempfile.mkstemp(suffix='.xlsx', dir=self.path.parent)
                    os.close(fd)
                    try:
                        book.save(temp)
                        with open(temp, 'rb+') as handle:
                            os.fsync(handle.fileno())
                        # Retry on Windows if file is temporarily locked by external reader
                        replaced = False
                        for delay in (0.05, 0.1, 0.2):
                            try:
                                os.replace(temp, self.path)
                                replaced = True
                                break
                            except PermissionError:
                                time.sleep(delay)
                        if not replaced:
                            os.replace(temp, self.path)
                    finally:
                        if os.path.exists(temp):
                            os.unlink(temp)
                return result
            finally:
                book.close()

    @staticmethod
    def _build_default_state(sheet, row, serial):
        reg_time = str(sheet.cell(row, 2).value or '')
        reg_bat = sheet.cell(row, 3).value
        try:
            reg_bat = int(reg_bat) if reg_bat is not None else 100
        except (ValueError, TypeError):
            reg_bat = 100
        status_val = str(sheet.cell(row, 14).value or '').strip()
        if not status_val:
            status_val = 'READY_FOR_AGING' if reg_bat == 100 else 'WAITING_FOR_100_PERCENT_CHARGE'
        highest_cp = 0
        for cp in range(1, 5):
            if sheet.cell(row, 2 * cp + 2).value is not None:
                highest_cp = cp
        next_cp = min(highest_cp + 1, 4)
        if status_val.startswith('AGING_HOUR_'):
            try:
                next_cp = int(status_val.split('_')[-1])
            except ValueError:
                pass
        obs, power = extract_observations_from_row(sheet, row)
        return {
            'serial_number': serial,
            'status': status_val,
            'pending_restart': None,
            'next_checkpoint': next_cp,
            'aging_started': reg_time or None,
            'next_due': None,
            'events': [{'action': 'excel_sync', 'battery': reg_bat, 'device_time': None, 'server_received': reg_time}],
            'registration_device_time': None,
            'last_server_received': reg_time,
            'last_device_time': None,
            'last_battery': reg_bat,
            'observations': obs,
            'power_test_result': power
        }

    @staticmethod
    def locate(book, serial):
        if not serial:
            return None
        target = str(serial).strip()
        for row in book['Devices'].iter_rows(min_row=2):
            val = row[0].value
            if val is not None and str(val).strip() == target:
                dev_row = row[0].row
                for meta in book['Workflow'].iter_rows(min_row=2):
                    m_val = meta[0].value
                    if m_val is not None and str(m_val).strip() == target:
                        raw_state = meta[1].value
                        state = json.loads(raw_state) if raw_state else {}
                        return dev_row, meta[0].row, state
                # Auto-create workflow state if missing (e.g. manually added in Excel)
                state = Store._build_default_state(book['Devices'], dev_row, target)
                book['Workflow'].append([target, json.dumps(state)])
                return dev_row, book['Workflow'].max_row, state
        return None
