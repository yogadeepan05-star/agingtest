import json
from openpyxl import load_workbook
from datetime import datetime, timezone, timedelta

path = '../data/aging_test.xlsx'
book = load_workbook(path)
ws = book['Workflow']
patched = []

for row in ws.iter_rows(min_row=2):
    serial = row[0].value
    if not serial:
        continue
    state = json.loads(row[1].value)
    if state.get('next_due'):
        old = state['next_due']
        # Set next_due to 2 minutes ago so checkpoint is immediately due
        state['next_due'] = (datetime.now(timezone.utc) - timedelta(minutes=2)).isoformat()
        row[1].value = json.dumps(state)
        patched.append('  ' + serial + ': ' + old + ' -> ' + state['next_due'])
    else:
        patched.append('  ' + serial + ': next_due is None, skipped')

book.save(path)
book.close()
print('Patched devices:')
for p in patched:
    print(p)
print('Done. Restart the server now.')
