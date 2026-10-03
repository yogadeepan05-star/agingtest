from concurrent.futures import ThreadPoolExecutor
import pytest
from fastapi.testclient import TestClient as FastAPITestClient

def TestClient(app):
    return FastAPITestClient(app, base_url="http://localhost")
from openpyxl import load_workbook
from app.main import create_app

SERIAL = 'T130R4CIK54677'
OTHER = 'T130R4CIK54678'

@pytest.fixture
def client(tmp_path):
    return TestClient(create_app(tmp_path / 'test.xlsx', interval=0))

def submit(client, action, battery=100, serial=SERIAL, target=None, **extra):
    target = None if action == 'register' else (target or serial)
    ticket = client.post('/api/captures', json={'action': action, 'serial_number': target}).json()['capture_token']
    body = dict(serial_number=serial, battery_percent=battery, device_timestamp='12:44 PM', capture_token=ticket, **extra)
    route = '/api/devices/register' if action == 'register' else f'/api/devices/{target}/' + ('aging/' + action if action.startswith('h') else action)
    return client.post(route, json=body)

def complete(client):
    assert submit(client, 'register').status_code == 200
    assert submit(client, 'start-aging').status_code == 200
    for n in range(1, 5):
        assert submit(client, f'h{n}', 90-n).status_code == 200
        assert client.post(f'/api/devices/{SERIAL}/restart', json={'checkpoint': n, 'confirmed': True}).status_code == 200

def test_registration_and_duplicate(client):
    assert submit(client, 'register', 84).json()['status'] == 'WAITING_FOR_100_PERCENT_CHARGE'
    assert submit(client, 'register').status_code == 409
    assert submit(client, 'start-aging', 99).status_code == 409
    assert submit(client, 'start-aging').json()['status'] == 'AGING_HOUR_1'

@pytest.mark.parametrize('value', [-1, 101, 4.2, True, '84'])
def test_invalid_battery(client, value):
    assert submit(client, 'register', value).status_code == 422

@pytest.mark.parametrize('serial', ['', 'BAD', '=SUM(1)', 'T130R4CIK5467!'])
def test_invalid_serial(client, serial):
    assert submit(client, 'register', serial=serial).status_code == 422

@pytest.mark.parametrize('battery,status', [(65, 'POST_AGING_CHARGE'), (70, 'PACKING_READY'), (100, 'PACKING_READY')])
def test_full_flow(client, battery, status):
    complete(client)
    assert submit(client, 'post-aging', battery).json()['status'] == status
    path = client.app.state.workflow.store.path
    book = load_workbook(path)
    assert book['Devices'].max_row == 2
    assert book['Devices']['M2'].value == '12:44 PM'
    assert book['Devices']['N2'].value == status
    book.close()
    assert TestClient(create_app(path)).get(f'/api/devices/{SERIAL}').json()['status'] == status

def test_wrong_order_and_restart(client):
    submit(client, 'register')
    assert submit(client, 'h2').status_code == 409
    submit(client, 'start-aging')
    for checkpoint in ['h2', 'h3', 'h4']:
        assert submit(client, checkpoint).status_code == 409
    submit(client, 'h1')
    assert submit(client, 'h2').status_code == 409
    assert client.post(f'/api/devices/{SERIAL}/restart', json={'checkpoint': 2, 'confirmed': True}).status_code == 409
    assert submit(client, 'post-aging').status_code == 409

def test_serial_mismatch_does_not_write(client):
    submit(client, 'register')
    before = client.get(f'/api/devices/{SERIAL}').json()
    assert submit(client, 'start-aging', serial=OTHER, target=SERIAL).status_code == 409
    assert client.get(f'/api/devices/{SERIAL}').json() == before

def test_replay(client):
    token = client.post('/api/captures', json={'action': 'register'}).json()['capture_token']
    body = {'serial_number': SERIAL, 'battery_percent': 100, 'capture_token': token}
    assert client.post('/api/devices/register', json=body).status_code == 200
    body['serial_number'] = OTHER
    assert client.post('/api/devices/register', json=body).status_code == 409

def test_hour_enforced(tmp_path):
    client = TestClient(create_app(tmp_path/'timed.xlsx', interval=3600))
    submit(client, 'register')
    submit(client, 'start-aging')
    assert submit(client, 'h1').status_code == 409

def test_concurrent_registration(client):
    with ThreadPoolExecutor(max_workers=4) as pool:
        results = list(pool.map(lambda _: submit(client, 'register').status_code, range(8)))
    assert results.count(200) == 1
    assert results.count(409) == 7

def test_forbid_image_and_large_body(client):
    assert submit(client, 'register', image='data:image/png;base64,123').status_code == 422
    assert client.post('/api/devices/register', content=b'x', headers={'Content-Type':'image/png'}).status_code == 415
    assert client.post('/api/devices/register', content=b'x'*5000, headers={'Content-Type':'application/json'}).status_code == 413

def test_post_charge_retry(client):
    complete(client)
    assert submit(client, 'post-aging', 65).status_code == 200
    assert submit(client, 'post-aging', 101).status_code == 422
    assert submit(client, 'post-aging', 70).json()['status'] == 'PACKING_READY'

def test_health_and_missing(client):
    assert client.get('/api/health').json() == {'status': 'ok'}
    assert client.get('/api/devices/missing').status_code == 404

def test_expired_capture(client):
    token = client.post('/api/captures', json={'action': 'register'}).json()['capture_token']
    client.app.state.workflow.store.transaction(lambda book: setattr(book['Captures']['D2'], 'value', '2000-01-01T00:00:00+00:00'))
    response = client.post('/api/devices/register', json={'serial_number': SERIAL, 'battery_percent':100, 'capture_token':token})
    assert response.status_code == 409
    assert client.get(f'/api/devices/{SERIAL}').status_code == 404

def test_missing_device_time_kept_separate(client):
    token = client.post('/api/captures', json={'action': 'register'}).json()['capture_token']
    response = client.post('/api/devices/register', json={'serial_number':SERIAL, 'battery_percent':84, 'capture_token':token})
    assert response.json()['last_device_time'] is None
    assert response.json()['last_server_received']

def test_existing_invalid_workbook_preserved(tmp_path):
    from openpyxl import Workbook
    path = tmp_path / 'existing.xlsx'
    book = Workbook()
    book.active['A1'] = 'Existing unrelated data'
    book.save(path)
    before = path.read_bytes()
    client = TestClient(create_app(path))
    assert client.get('/api/health').status_code == 503
    assert path.read_bytes() == before


def test_checkpoint_observations_and_power_test(tmp_path):
    path = tmp_path / 'obs.xlsx'
    client = TestClient(create_app(path, interval=0))

    # 1. Register device
    assert submit(client, 'register').status_code == 200
    assert submit(client, 'start-aging').status_code == 200

    # 2. Checkpoint H1: No issue, with optional remark
    res_h1 = submit(client, 'h1', 89, has_issue='no', remarks='All normal at H1')
    assert res_h1.status_code == 200
    assert res_h1.json()['observations']['h1']['has_issue'] == 'no'
    assert client.post(f'/api/devices/{SERIAL}/restart', json={'checkpoint': 1, 'confirmed': True}).status_code == 200

    # 3. Checkpoint H2: Issue with single category
    res_h2 = submit(client, 'h2', 88, has_issue='yes', issue_categories=['Display issue'], remarks='Minor display flicker')
    assert res_h2.status_code == 200
    assert res_h2.json()['observations']['h2']['has_issue'] == 'yes'
    assert res_h2.json()['observations']['h2']['categories'] == ['Display issue']
    # Verify earlier checkpoint H1 observation is NOT overwritten
    assert res_h2.json()['observations']['h1']['has_issue'] == 'no'
    assert client.post(f'/api/devices/{SERIAL}/restart', json={'checkpoint': 2, 'confirmed': True}).status_code == 200

    # 4. Checkpoint H3: Issue with multiple categories and formula injection test
    res_h3 = submit(
        client, 'h3', 87,
        has_issue='yes',
        issue_categories=['Crashing / hanging issue', 'Other issue'],
        remarks='=HYPERLINK("http://malicious.com")'
    )
    assert res_h3.status_code == 200
    assert res_h3.json()['observations']['h3']['has_issue'] == 'yes'
    assert 'Crashing / hanging issue' in res_h3.json()['observations']['h3']['categories']
    assert 'Other issue' in res_h3.json()['observations']['h3']['categories']
    assert client.post(f'/api/devices/{SERIAL}/restart', json={'checkpoint': 3, 'confirmed': True}).status_code == 200

    # 5. Checkpoint H4: No issue, no remarks
    res_h4 = submit(client, 'h4', 86, has_issue='no')
    assert res_h4.status_code == 200
    assert res_h4.json()['observations']['h4']['has_issue'] == 'no'
    assert client.post(f'/api/devices/{SERIAL}/restart', json={'checkpoint': 4, 'confirmed': True}).status_code == 200

    # 6. Post-aging: Issue, remarks, and power test Pass
    res_post = submit(
        client, 'post-aging', 85,
        has_issue='yes',
        issue_categories=['Display issue'],
        remarks='Final packing inspection remark',
        power_test_result='Pass'
    )
    assert res_post.status_code == 200
    assert res_post.json()['status'] == 'PACKING_READY'
    assert res_post.json()['power_test_result'] == 'Pass'
    assert res_post.json()['observations']['post']['has_issue'] == 'yes'

    # 7. Check Excel direct cell storage (Columns 15–30)
    book = load_workbook(path)
    sheet = book['Devices']
    assert sheet.max_row == 2
    # H1
    assert sheet.cell(2, 15).value == 'No'
    assert sheet.cell(2, 16).value is None
    assert sheet.cell(2, 17).value == 'All normal at H1'
    # H2
    assert sheet.cell(2, 18).value == 'Yes'
    assert sheet.cell(2, 19).value == 'Display issue'
    assert sheet.cell(2, 20).value == 'Minor display flicker'
    # H3
    assert sheet.cell(2, 21).value == 'Yes'
    assert sheet.cell(2, 22).value == 'Crashing / hanging issue, Other issue'
    # Formula injection sanitized with leading single quote:
    assert sheet.cell(2, 23).value == "'=HYPERLINK(\"http://malicious.com\")"
    # H4
    assert sheet.cell(2, 24).value == 'No'
    assert sheet.cell(2, 25).value is None
    # Post
    assert sheet.cell(2, 27).value == 'Yes'
    assert sheet.cell(2, 28).value == 'Display issue'
    assert sheet.cell(2, 29).value == 'Final packing inspection remark'
    # Power test
    assert sheet.cell(2, 30).value == 'Pass'
    book.close()

    # 8. Reload app and verify all fields are preserved
    reloaded_client = TestClient(create_app(path))
    dev = reloaded_client.get(f'/api/devices/{SERIAL}').json()
    assert dev['observations']['h1']['has_issue'] == 'no'
    assert dev['observations']['h2']['has_issue'] == 'yes'
    assert dev['observations']['h3']['has_issue'] == 'yes'
    assert dev['observations']['h4']['has_issue'] == 'no'
    assert dev['observations']['post']['has_issue'] == 'yes'
    assert dev['power_test_result'] == 'Pass'


def test_historical_records_preserved_as_blank_without_defaulting(tmp_path):
    from openpyxl import Workbook
    path = tmp_path / 'legacy.xlsx'
    book = Workbook()
    sheet = book.active
    sheet.title = 'Devices'
    # 14 base headers only
    headers = ['Serial Number', 'Device Registration Time', 'Registration Battery %',
               'H1 Battery %', 'H1 Timestamp', 'H2 Battery %', 'H2 Timestamp',
               'H3 Battery %', 'H3 Timestamp', 'H4 Battery %', 'H4 Timestamp',
               'Post-Aging Battery %', 'Post-Aging Timestamp', 'Final Status']
    sheet.append(headers)
    sheet.append([SERIAL, '2026-09-01T10:00:00', 100, 95, '11:00 AM', None, None, None, None, None, None, None, None, 'AGING_HOUR_2'])
    book.save(path)
    book.close()

    client = TestClient(create_app(path))
    dev = client.get(f'/api/devices/{SERIAL}').json()
    # Must preserve missing new fields as blank/None, NEVER auto-label as No or Pass!
    assert dev['observations']['h1'] is None
    assert dev['observations']['h2'] is None
    assert dev['observations']['post'] is None
    assert dev['power_test_result'] is None


def test_observation_validation_rules(client):
    assert submit(client, 'register').status_code == 200
    assert submit(client, 'start-aging').status_code == 200

    # Yes without issue categories must be rejected (422)
    assert submit(client, 'h1', 90, has_issue='yes', issue_categories=[]).status_code == 422
    assert submit(client, 'h1', 90, has_issue='yes').status_code == 422

    # No with issue categories must be rejected (422)
    assert submit(client, 'h1', 90, has_issue='no', issue_categories=['Display issue']).status_code == 422

    # Invalid category must be rejected (422)
    assert submit(client, 'h1', 90, has_issue='yes', issue_categories=['Random issue']).status_code == 422

    # Invalid power test must be rejected (422)
    assert submit(client, 'h1', 90, power_test_result='Maybe').status_code == 422

