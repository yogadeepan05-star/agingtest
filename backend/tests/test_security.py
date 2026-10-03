import json
import subprocess
import sys
from concurrent.futures import ThreadPoolExecutor
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from app.main import create_app
from test_workflow import submit, SERIAL

@pytest.fixture
def client(tmp_path):
    return TestClient(create_app(tmp_path/'secure.xlsx', interval=0), base_url='http://localhost')

BAD_SERIALS=['=1+1','+SUM(1)','-1','@SUM(1)','../../test',r'..\..\test','/etc/passwd',r'C:\Windows\System32','<script>alert(1)</script>',"' OR 1=1 --",'T130R4CIK\r\n54677','设备123','x'*65,'']

@pytest.mark.parametrize('serial',BAD_SERIALS)
def test_capture_serial_cannot_write_formulas_or_paths(client,serial):
    response=client.post('/api/captures',json={'action':'h1','serial_number':serial})
    assert response.status_code==422
    assert not client.app.state.workflow.store.path.exists()

@pytest.mark.parametrize('serial',BAD_SERIALS)
def test_reading_serial_rejected(client,serial):
    response=submit(client,'register',serial=serial)
    assert response.status_code==422
    book=load_workbook(client.app.state.workflow.store.path)
    assert book['Devices'].max_row==1
    assert all(cell.data_type!='f' for sheet in book for row in sheet for cell in row)
    book.close()

@pytest.mark.parametrize('value',['25:99 PM','<script>','12:44 PM\n=1+1','x'*100,5,True])
def test_invalid_time(client,value):
    token=client.post('/api/captures',json={'action':'register'}).json()['capture_token']
    response=client.post('/api/devices/register',json={'serial_number':SERIAL,'battery_percent':84,'device_timestamp':value,'capture_token':token})
    assert response.status_code==422

@pytest.mark.parametrize('payload',['{','[]','null','{"action":"register","action":"h1"}','{"battery_percent":NaN}','{"a":Infinity}'])
def test_malformed_json(client,payload):
    response=client.post('/api/captures',content=payload,headers={'Content-Type':'application/json'})
    assert response.status_code==422
    assert response.json() == {"detail": "Invalid JSON object."}
    assert 'no-store' in response.headers['cache-control']

def test_cross_site_host_and_docs(client):
    assert client.post('/api/captures',json={'action':'register'},headers={'Origin':'https://attacker.example'}).status_code==403
    assert client.post('/api/captures',json={'action':'register'},headers={'Sec-Fetch-Site':'cross-site'}).status_code==403
    assert client.get('/api/health',headers={'Host':'attacker.example'}).status_code==400
    response=client.get('/api/health',headers={'Origin':'https://localhost:5173'})
    assert response.status_code==200
    assert 'access-control-allow-origin' not in response.headers
    assert client.get('/docs').status_code==404
    assert client.get('/openapi.json').status_code==404

def test_bounded_rate(tmp_path):
    c=TestClient(create_app(tmp_path/'rate.xlsx',rate_limit=2),base_url='http://localhost')
    assert c.get('/api/health').status_code==200
    assert c.get('/api/health').status_code==200
    assert c.get('/api/health').status_code==429

def test_missing_or_extra_capture_fields(client):
    for payload in [{},{'action':'h1'},{'action':'register','serial_number':SERIAL},{'action':'register','image':'abc'},{'action':'unknown'}]:
        assert client.post('/api/captures',json=payload).status_code==422

def test_stale_ticket_after_other_operator_updates_state(client):
    submit(client,'register');submit(client,'start-aging')
    old=client.post('/api/captures',json={'action':'h2','serial_number':SERIAL}).json()['capture_token']
    submit(client,'h1',90)
    client.post(f'/api/devices/{SERIAL}/restart',json={'checkpoint':1,'confirmed':True})
    response=client.post(f'/api/devices/{SERIAL}/aging/h2',json={'serial_number':SERIAL,'battery_percent':85,'capture_token':old})
    assert response.status_code==409
    assert client.get(f'/api/devices/{SERIAL}').json()['values'][5] is None

def test_concurrent_checkpoint_cannot_overwrite(client):
    submit(client,'register');submit(client,'start-aging')
    with ThreadPoolExecutor(max_workers=4) as pool:
        results=list(pool.map(lambda _:submit(client,'h1',90).status_code,range(4)))
    assert sorted(results)==[200,409,409,409]
    state=client.get(f'/api/devices/{SERIAL}').json()
    assert state['pending_restart']==1 and state['values'][3]==90

@pytest.mark.parametrize('confirmation',[False,1,'true',None])
def test_restart_requires_real_boolean(client,confirmation):
    submit(client,'register');submit(client,'start-aging');submit(client,'h1')
    response=client.post(f'/api/devices/{SERIAL}/restart',json={'checkpoint':1,'confirmed':confirmation})
    assert response.status_code==422

def test_atomic_replace_failure_preserves_workbook(client,monkeypatch):
    submit(client,'register')
    path=client.app.state.workflow.store.path;before=path.read_bytes()
    def denied(*_args):raise PermissionError('private path should not leak')
    monkeypatch.setattr('app.storage.os.replace',denied)
    response=client.post('/api/captures',json={'action':'start-aging','serial_number':SERIAL})
    assert response.status_code==503
    assert 'private path' not in response.text
    assert path.read_bytes()==before
    assert not list(path.parent.glob('tmp*.xlsx'))

def test_process_lock_prevents_duplicate_rows(tmp_path):
    path=tmp_path/'processes.xlsx'
    script='''import sys
from fastapi.testclient import TestClient
from app.main import create_app
c=TestClient(create_app(sys.argv[1]),base_url='http://localhost')
r=c.post('/api/captures',json={'action':'register'})
assert r.status_code==200, r.text
t=r.json()['capture_token']
r=c.post('/api/devices/register',json={'serial_number':'T130R4CIK54677','battery_percent':84,'capture_token':t})
print(r.status_code)
'''
    processes=[subprocess.Popen([sys.executable,'-c',script,str(path)],stdout=subprocess.PIPE,stderr=subprocess.PIPE,text=True) for _ in range(4)]
    statuses=[]
    for process in processes:
        stdout,stderr=process.communicate(timeout=30)
        assert process.returncode==0,stderr
        statuses.append(int(stdout.strip()))
    assert sorted(statuses)==[200,409,409,409]
    book=load_workbook(path);assert book['Devices'].max_row==2;book.close()
